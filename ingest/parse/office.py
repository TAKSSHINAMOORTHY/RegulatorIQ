"""Parsers for Office formats (.docx and .xlsx)."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import List

import docx
import openpyxl

from ingest.parse.models import ParsedBlock, ParsedDocument, ParsedPage, ParsedTable
from ingest.parse.pdf import table_to_markdown

logger = logging.getLogger(__name__)


def parse_docx(file_path: Path, doc_id: str) -> ParsedDocument:
    """Parse Microsoft Word (.docx) circulars."""
    doc = docx.Document(str(file_path))
    paras: List[str] = []
    blocks: List[ParsedBlock] = []
    tables: List[ParsedTable] = []

    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paras.append(txt)
            blocks.append(ParsedBlock(text=txt, type="text"))

    # Extract docx tables
    for idx, tbl in enumerate(doc.tables, start=1):
        raw_rows: List[List[str]] = []
        for row in tbl.rows:
            raw_rows.append([cell.text.strip() for cell in row.cells])
        md = table_to_markdown(raw_rows)
        if md:
            tables.append(ParsedTable(page_no=1, markdown=md))

    full_text = "\n\n".join(paras)
    return ParsedDocument(
        doc_id=doc_id,
        pages=[ParsedPage(page_no=1, text=full_text, blocks=blocks)],
        full_text=full_text,
        tables=tables,
        extraction_method="docx",
    )


def parse_xlsx(file_path: Path, doc_id: str) -> ParsedDocument:
    """Parse Microsoft Excel (.xlsx) returns or data sheets."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    tables: List[ParsedTable] = []
    sheet_texts: List[str] = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        raw_rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            if any(row):
                raw_rows.append([str(v or "").strip() for v in row])
        md = table_to_markdown(raw_rows)
        if md:
            tables.append(
                ParsedTable(page_no=sheet_idx, markdown=md, caption=f"Sheet: {sheet_name}")
            )
            sheet_texts.append(f"### Sheet: {sheet_name}\n\n" + md)

    full_text = "\n\n".join(sheet_texts)
    return ParsedDocument(
        doc_id=doc_id,
        pages=[ParsedPage(page_no=1, text=full_text, blocks=[ParsedBlock(text=full_text)])],
        full_text=full_text,
        tables=tables,
        extraction_method="xlsx",
    )
